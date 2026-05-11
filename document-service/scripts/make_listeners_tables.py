#
# НЕ ТРОГАТЬ!!!!!!!!!!
#

import sys
import json
from docx import Document

def add_data_to_table(path_to_docx, path_to_raw_docx, json_data):
    doc = Document(path_to_raw_docx)

    data = json.loads(json_data)

    found_table = doc.tables[1]

    i = 0
    for listener in data:

        i = i + 1

        row_cells = found_table.add_row().cells
        row_cells[0].text = str(i)
        row_cells[1].text = listener.get("fio", "")
        row_cells[2].text = listener.get("date_birth", "")
        row_cells[3].text = listener.get("document", "")
        row_cells[4].text = listener.get("SNILS", "")
        row_cells[5].text = listener.get("email", "")

    found_table = doc.tables[2]
    i = 0
    for listener in data:

        i = i + 1
        row_cells = found_table.add_row().cells
        row_cells[0].text = str(i)
        row_cells[1].text = listener.get("fio", "")
        row_cells[2].text = listener.get("period", "")
        row_cells[3].text = listener.get("obem", "")
        row_cells[4].text = listener.get("cost", "")

    doc.save(path_to_docx)

if __name__ == "__main__":
    path = sys.argv[1]
    path_to_raw_docx = sys.argv[2]
    json_data = sys.argv[3]

    add_data_to_table(path, path_to_raw_docx, json_data)        